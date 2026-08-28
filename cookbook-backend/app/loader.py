"""Turn an uploaded file of (almost) any type into plain text.

Dispatch is by extension. Anything without a dedicated extractor is decoded as
text; anything that is not text is salvaged for readable strings; images and
image-only PDFs are transcribed by Gemini (see `app.ocr`). Only media files and
archives are turned away outright, because there is no document inside them to
index.
"""

import csv
import io
import json
import logging
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from pypdf import PdfReader

from app.config import settings

logger = logging.getLogger(__name__)


class UnsupportedDocumentError(ValueError):
    """Raised when a file's format cannot be turned into plain text."""


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------

# Runs of printable bytes, and the same again as UTF-16LE (ASCII interleaved
# with NULs) - between them they cover how legacy Office formats store text.
_ASCII_RUN = re.compile(rb"[\x09\x0a\x0d\x20-\x7e\xa0-\xff]{4,}")
_UTF16_RUN = re.compile(rb"(?:[\x09\x0a\x0d\x20-\x7e]\x00){4,}")

_WORDLIKE_RE = re.compile(r"[A-Za-z0-9\s.,;:!?'\"()\-]")


def _read_text(file_path: str) -> str:
    data = Path(file_path).read_bytes()
    if b"\x00" in data[:4096]:
        raise UnsupportedDocumentError(
            "This file looks like binary data rather than a readable document."
        )
    return _decode(data)


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _normalise_newlines(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _looks_like_text(text: str) -> bool:
    """Guard against indexing binary noise that happens to decode."""
    if len(text) < 40:
        return False
    wordlike = len(_WORDLIKE_RE.findall(text))
    return wordlike / len(text) > 0.75


def _is_prose(run: str) -> bool:
    """Is this run of characters sentence-like, or just binary that decoded?

    Judged per run, not over the whole file: a legacy .doc interleaves its body
    text with structural noise, and scoring them together buries the prose.
    """
    run = run.strip()
    if len(run) < 8 or not re.search(r"[A-Za-z]{3}", run):
        return False
    return len(_WORDLIKE_RE.findall(run)) / len(run) > 0.8


def _salvage_text(data: bytes) -> str:
    """Pull readable strings out of a binary blob, format-agnostically.

    Legacy Office files store their body text as long contiguous runs, so this
    recovers the prose (alongside the occasional field code).
    """
    runs = [_decode(match.group()) for match in _ASCII_RUN.finditer(data)]
    runs += [match.group().decode("utf-16-le", errors="replace") for match in _UTF16_RUN.finditer(data)]
    return "\n".join(run.strip() for run in runs if _is_prose(run))


def _strip_xml(markup: str) -> str:
    """Text content of an XML/HTML document, tags and entities resolved."""
    parser = _TextExtractor()
    parser.feed(markup)
    parser.close()
    return "\n".join(parser.parts)


class _TextExtractor(HTMLParser):
    _SKIP = {"script", "style", "head"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skipping = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skipping += 1

    def handle_endtag(self, tag):
        if tag in self._SKIP and self._skipping:
            self._skipping -= 1

    def handle_data(self, data):
        if not self._skipping and data.strip():
            self.parts.append(data.strip())


# --------------------------------------------------------------------------
# loaders
# --------------------------------------------------------------------------


def load_pdf(file_path: str) -> str:
    """Extract embedded text; transcribe the pages if the PDF is a scan."""
    reader = PdfReader(file_path)
    text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if text or not settings.ocr_enabled:
        return text

    logger.info("No embedded text in %s; transcribing it instead", file_path)
    return _transcribe(file_path, "application/pdf")


def _transcribe(file_path: str, mime_type: str) -> str:
    """Hand a file to Gemini for transcription, as a readable failure."""
    from app.ocr import transcribe

    try:
        return transcribe(Path(file_path).read_bytes(), mime_type)
    except ValueError as exc:  # size limit - the message is meant for the user
        raise UnsupportedDocumentError(str(exc)) from exc


def load_image(file_path: str) -> str:
    from app.ocr import IMAGE_MIME_TYPES

    if not settings.ocr_enabled:
        raise UnsupportedDocumentError(
            "Image transcription is turned off on this server, so images cannot be read."
        )
    suffix = Path(file_path).suffix.lower()
    text = _transcribe(file_path, IMAGE_MIME_TYPES[suffix])
    if not text:
        raise UnsupportedDocumentError("No legible text was found in this image.")
    return text


def load_docx(file_path: str) -> str:
    from docx import Document

    document = Document(file_path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(part for part in parts if part.strip())


def load_pptx(file_path: str) -> str:
    from pptx import Presentation

    parts = []
    for index, slide in enumerate(Presentation(file_path).slides, start=1):
        parts.append(f"[Slide {index}]")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(part for part in parts if part.strip())


def load_xlsx(file_path: str) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    try:
        for sheet in workbook.worksheets:
            parts.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                line = " | ".join("" if value is None else str(value) for value in row)
                if line.strip(" |"):
                    parts.append(line)
    finally:
        workbook.close()
    return "\n".join(parts)


def load_xls(file_path: str) -> str:
    """Legacy Excel (BIFF). xlrd 2.x reads .xls and nothing else."""
    import xlrd

    book = xlrd.open_workbook(file_path)
    parts = []
    for sheet in book.sheets():
        parts.append(f"[Sheet: {sheet.name}]")
        for index in range(sheet.nrows):
            line = " | ".join(str(cell.value) for cell in sheet.row(index))
            if line.strip(" |"):
                parts.append(line)
    return "\n".join(parts)


# Streams holding the body text of the legacy binary Office formats.
_OLE_TEXT_STREAMS = ("WordDocument", "PowerPoint Document", "Book", "Workbook")

# PowerPoint record types carrying text, and the master-slide prompts that are
# boilerplate rather than content.
_PPT_TEXT_CHARS = 0x0FA0  # UTF-16LE
_PPT_TEXT_BYTES = 0x0FA8  # cp1252
_PPT_PLACEHOLDERS = ("Click to edit", "Click to add")
# Outline levels from the slide master template, not slide content.
_PPT_OUTLINE_RE = re.compile(r"^(second|third|fourth|fifth) level$", re.IGNORECASE)


def _ppt_records(data: bytes, start: int = 0, end: int | None = None):
    """Walk PowerPoint's record tree, descending into container records."""
    end = len(data) if end is None else end
    position = start
    while position + 8 <= end:
        version = int.from_bytes(data[position : position + 2], "little") & 0x000F
        record_type = int.from_bytes(data[position + 2 : position + 4], "little")
        length = int.from_bytes(data[position + 4 : position + 8], "little")
        body_start = position + 8
        body_end = min(body_start + length, end)
        if version == 0x0F:
            yield from _ppt_records(data, body_start, body_end)
        else:
            yield record_type, data[body_start:body_end]
        position = max(body_end, body_start)


def _ppt_text(data: bytes) -> str:
    parts = []
    for record_type, body in _ppt_records(data):
        if record_type == _PPT_TEXT_CHARS:
            parts.append(body.decode("utf-16-le", errors="replace"))
        elif record_type == _PPT_TEXT_BYTES:
            parts.append(body.decode("cp1252", errors="replace"))
    lines = []
    for part in parts:
        for line in part.replace("\r", "\n").replace("\x0b", "\n").split("\n"):
            line = line.strip()
            if len(line) < 2 or line.startswith(_PPT_PLACEHOLDERS) or _PPT_OUTLINE_RE.match(line):
                continue
            lines.append(line)
    return "\n".join(dict.fromkeys(lines))


def load_ole(file_path: str) -> str:
    """Legacy binary Office (.doc, .ppt): best-effort text salvage.

    These formats have no maintained Python parser, so the body text is
    recovered from the document stream rather than parsed structurally. Expect
    the prose, plus the occasional field code.
    """
    import olefile

    if not olefile.isOleFile(file_path):
        raise UnsupportedDocumentError(
            "This file is not a readable Office document."
        )
    with olefile.OleFileIO(file_path) as ole:
        entries = ["/".join(entry) for entry in ole.listdir()]
        streams = [name for name in _OLE_TEXT_STREAMS if name in entries] or entries
        parts = []
        for name in streams:
            data = ole.openstream(name).read()
            # PowerPoint keeps text in typed records; reading them avoids the
            # embedded-XML noise that a raw salvage would pick up.
            parsed = _ppt_text(data) if name == "PowerPoint Document" else ""
            parts.append(parsed or _salvage_text(data))
        text = "\n".join(parts)

    if not _looks_like_text(text):
        raise UnsupportedDocumentError(
            "No readable text could be recovered from this file. Re-saving it in a "
            "modern format (.docx, .pptx) will give a much better result."
        )
    return text


def load_zip_documents(file_path: str, members: tuple[str, ...] = (), suffixes: tuple[str, ...] = ()) -> str:
    """Text of the XML/XHTML parts inside a zip-backed document (EPUB, ODF)."""
    with zipfile.ZipFile(file_path) as archive:
        names = [name for name in archive.namelist() if name in members] or sorted(
            name for name in archive.namelist() if name.lower().endswith(suffixes)
        )
        return "\n".join(_strip_xml(_decode(archive.read(name))) for name in names)


def load_epub(file_path: str) -> str:
    return load_zip_documents(file_path, suffixes=(".xhtml", ".html", ".htm"))


def load_opendocument(file_path: str) -> str:
    return load_zip_documents(file_path, members=("content.xml",))


def load_rtf(file_path: str) -> str:
    return _rtf_to_text(_read_text(file_path))


# Control words whose group holds metadata or binary data, not body text.
_RTF_SKIP_GROUPS = {
    "fonttbl", "colortbl", "stylesheet", "listtable", "listoverridetable",
    "info", "pict", "object", "themedata", "colorschememapping", "datastore",
    "latentstyles", "generator", "xmlnstbl", "rsidtbl", "mmathPr",
}
# Control words that stand for a character or a break.
_RTF_LITERALS = {
    "par": "\n", "line": "\n", "tab": "\t", "sect": "\n\n", "page": "\n\n",
    "emdash": "-", "endash": "-", "lquote": "'", "rquote": "'",
    "ldblquote": '"', "rdblquote": '"', "bullet": "* ", "nbsp": " ",
}
_RTF_TOKEN = re.compile(
    r"\\([a-zA-Z]{1,32})(-?\d{1,10})?[ ]?|\\'([0-9a-fA-F]{2})|\\([^a-zA-Z])|([{}])|([^\\{}]+)"
)


def _rtf_to_text(rtf: str) -> str:
    """Strip RTF control words, keeping the body text."""
    out: list[str] = []
    depth = 0
    skip_until: int | None = None
    # \uN is followed by \ucN replacement characters for readers that cannot do
    # Unicode; they must be dropped or every smart quote leaves a stray "?".
    unicode_skip = 1
    pending_skip = 0
    for match in _RTF_TOKEN.finditer(rtf):
        word, arg, hexchar, escaped, brace, plain = match.groups()
        if brace == "{":
            depth += 1
        elif brace == "}":
            depth -= 1
            if skip_until is not None and depth < skip_until:
                skip_until = None
        elif skip_until is not None:
            continue
        elif word:
            if word in _RTF_SKIP_GROUPS:
                skip_until = depth
            elif word == "uc" and arg:
                unicode_skip = max(int(arg), 0)
            elif word == "u" and arg:
                out.append(chr(int(arg) % 65536))
                pending_skip = unicode_skip
            elif word in _RTF_LITERALS:
                out.append(_RTF_LITERALS[word])
        elif hexchar:
            if pending_skip:
                pending_skip -= 1
            else:
                out.append(bytes([int(hexchar, 16)]).decode("cp1252", errors="replace"))
        elif escaped:
            if pending_skip:
                pending_skip -= 1
            else:
                out.append("" if escaped == "*" else escaped)
        elif plain:
            if pending_skip:
                dropped = min(pending_skip, len(plain))
                plain, pending_skip = plain[dropped:], pending_skip - dropped
            out.append(plain)
    text = "".join(out)
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())


def load_csv(file_path: str) -> str:
    text = _read_text(file_path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    rows = csv.reader(io.StringIO(text), dialect)
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows if any(row))


def load_html(file_path: str) -> str:
    return _strip_xml(_read_text(file_path))


def load_json(file_path: str) -> str:
    text = _read_text(file_path)
    try:
        return json.dumps(json.loads(text), indent=2, ensure_ascii=False)
    except json.JSONDecodeError:
        return text


def load_text(file_path: str) -> str:
    return _read_text(file_path)


def load_unknown(file_path: str) -> str:
    """Last resort: read it as text, and if it is binary, salvage what reads."""
    data = Path(file_path).read_bytes()
    if b"\x00" not in data[:4096]:
        return _decode(data)
    text = _salvage_text(data)
    if not _looks_like_text(text):
        raise UnsupportedDocumentError(
            "No readable text could be found in this file. Please upload a document."
        )
    logger.info("Salvaged %d chars of text from an unrecognised binary file", len(text))
    return text


# --------------------------------------------------------------------------
# dispatch
# --------------------------------------------------------------------------

LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".doc": load_ole,
    ".pptx": load_pptx,
    ".ppt": load_ole,
    ".xlsx": load_xlsx,
    ".xlsm": load_xlsx,
    ".xls": load_xls,
    ".odt": load_opendocument,
    ".ods": load_opendocument,
    ".odp": load_opendocument,
    ".epub": load_epub,
    ".rtf": load_rtf,
    ".csv": load_csv,
    ".tsv": load_csv,
    ".html": load_html,
    ".htm": load_html,
    ".xhtml": load_html,
    ".json": load_json,
}

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".text", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".xml", ".srt", ".vtt", ".tex", ".sql", ".py",
    ".js", ".ts", ".java", ".go", ".rb", ".sh",
}

IMAGE_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff",
    ".heic", ".heif",
}

# Files with no document inside them to index.
MEDIA_EXTENSIONS = {
    ".mp3", ".wav", ".m4a", ".aac", ".ogg", ".flac", ".wma",
    ".mp4", ".mov", ".avi", ".mkv", ".webm", ".wmv", ".m4v",
}
ARCHIVE_EXTENSIONS = {".zip", ".rar", ".7z", ".tar", ".gz", ".bz2", ".xz", ".iso"}

SUPPORTED_EXTENSIONS = sorted(LOADERS.keys() | TEXT_EXTENSIONS | IMAGE_EXTENSIONS)


def load_file(file_path: str, filename: str | None = None) -> str:
    """Extract plain text from a document of any type."""
    suffix = Path(filename or file_path).suffix.lower()

    if suffix in MEDIA_EXTENSIONS:
        raise UnsupportedDocumentError(
            "Audio and video files are not supported. Please upload a document."
        )
    if suffix in ARCHIVE_EXTENSIONS:
        raise UnsupportedDocumentError(
            "Archives are not supported. Please unpack it and upload the documents inside."
        )

    if suffix in IMAGE_EXTENSIONS:
        loader = load_image
    else:
        loader = LOADERS.get(suffix, load_text if suffix in TEXT_EXTENSIONS else load_unknown)

    try:
        # Loaders return whatever line endings their format uses, and legacy
        # Word separates paragraphs with a bare CR; the chunker splits on LF.
        return _normalise_newlines(loader(file_path))
    except UnsupportedDocumentError:
        raise
    except ImportError as exc:
        raise UnsupportedDocumentError(
            f"Support for {suffix or 'this type of'} file is not installed on the server."
        ) from exc
    except Exception as exc:
        raise UnsupportedDocumentError(
            f"The file could not be read as {suffix.lstrip('.').upper() or 'text'}. "
            "It may be corrupted or saved in a different format."
        ) from exc


def load_documents_dir(directory: str | None = None) -> list[dict]:
    docs_dir = Path(directory or settings.documents_dir)
    documents = []
    for path in sorted(docs_dir.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        try:
            text = load_file(str(path))
        except UnsupportedDocumentError:
            logger.warning("Skipping %s: no readable text", path.name)
            continue
        documents.append({"source": path.name, "text": text})
    return documents
